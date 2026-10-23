"""
Import d'articles et de stock depuis un fichier Excel (.xlsx) ou CSV.

Une ligne = une variante. Les articles sont regroupés par référence : plusieurs
lignes avec la même référence créent plusieurs variantes du même article.
L'import est « tout ou rien » : à la moindre erreur, rien n'est enregistré et
la liste des erreurs (avec numéro de ligne) est renvoyée.
"""

import csv
import io
from decimal import Decimal, InvalidOperation

from django.db import transaction

from apps.common.exceptions import BusinessError

from .models import AgeRange, Article, ArticleType, ArticleVariant, Category, Color, Gender, Size

COLUMNS = [
    ("reference", "Référence de l'article (obligatoire)", "ROB-001"),
    ("nom", "Nom de l'article (obligatoire)", "Robe à volants"),
    ("categorie", "Sous-catégorie, ou « Catégorie > Sous-catégorie » (obligatoire)", "Vêtements fille > Robes"),
    ("type", "Type : " + ", ".join(label for _, label in ArticleType.choices), "Vêtement"),
    ("genre", "Fille, Garçon ou Mixte", "Fille"),
    ("tranches_age", "Tranches d'âge séparées par |", "4–6 ans|6–8 ans"),
    ("marque", "Marque", ""),
    ("prix_achat", "Prix d'achat (obligatoire)", "4500"),
    ("prix_vente", "Prix de vente (obligatoire)", "9000"),
    ("taille", "Taille (créée si elle n'existe pas)", "6A"),
    ("couleur", "Couleur (créée si elle n'existe pas)", "Rose"),
    ("code_barres", "Code-barres", ""),
    ("quantite", "Quantité à ajouter au stock", "5"),
    ("seuil_alerte", "Seuil d'alerte stock faible", "2"),
    ("numero_lot", "Numéro de lot (obligatoire pour un produit réglementé)", ""),
]
REQUIRED = ["reference", "nom", "categorie", "prix_achat", "prix_vente"]


def _normalize(value):
    return str(value).strip() if value is not None else ""


def read_rows(uploaded_file):
    """Renvoie une liste de dictionnaires {colonne: valeur} à partir d'un .xlsx ou .csv."""
    name = uploaded_file.name.lower()
    if name.endswith(".xlsx"):
        from openpyxl import load_workbook

        sheet = load_workbook(uploaded_file, read_only=True, data_only=True).active
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [_normalize(h).lower() for h in rows[0]]
        return [dict(zip(headers, (_normalize(v) for v in row))) for row in rows[1:]
                if any(v not in (None, "") for v in row)]
    if name.endswith(".csv"):
        text = uploaded_file.read().decode("utf-8-sig")
        dialect = csv.Sniffer().sniff(text.splitlines()[0] if text else ";", delimiters=";,")
        reader = csv.DictReader(io.StringIO(text), dialect=dialect)
        return [{_normalize(k).lower(): _normalize(v) for k, v in row.items() if k}
                for row in reader if any(_normalize(v) for v in row.values())]
    raise BusinessError({"file": "Format non pris en charge. Utilisez un fichier .xlsx ou .csv."})


def _choice(value, choices, default):
    if not value:
        return default
    for key, label in choices:
        if value.lower() in (key.lower(), label.lower()):
            return key
    raise ValueError(f"valeur « {value} » inconnue")


def _decimal(value, field):
    try:
        number = Decimal(value.replace(" ", "").replace(",", "."))
    except (InvalidOperation, AttributeError):
        raise ValueError(f"{field} « {value} » n'est pas un nombre") from None
    if number < 0:
        raise ValueError(f"{field} ne peut pas être négatif")
    return number


def _integer(value, field, default=0):
    if not value:
        return default
    try:
        number = int(Decimal(value.replace(",", ".")))
    except InvalidOperation:
        raise ValueError(f"{field} « {value} » n'est pas un nombre entier") from None
    if number < 0:
        raise ValueError(f"{field} ne peut pas être négatif")
    return number


def _find_category(value):
    if ">" in value:
        parent_name, child_name = (p.strip() for p in value.split(">", 1))
        return Category.objects.get(name__iexact=child_name, parent__name__iexact=parent_name)
    matches = list(Category.objects.filter(name__iexact=value))
    if len(matches) > 1:
        raise ValueError(f"catégorie « {value} » ambiguë, écrivez « Catégorie > {value} »")
    if not matches:
        raise Category.DoesNotExist
    return matches[0]


def import_articles(uploaded_file, user, *, dry_run=False):
    from apps.inventory.models import StockMovement
    from apps.inventory.services import apply_movement

    rows = read_rows(uploaded_file)
    if not rows:
        raise BusinessError({"file": "Le fichier est vide."})
    missing = [c for c in REQUIRED if c not in rows[0]]
    if missing:
        raise BusinessError({"file": f"Colonnes manquantes : {', '.join(missing)}. Téléchargez le modèle."})

    errors = []
    summary = {"articles_created": 0, "articles_updated": 0, "variants_created": 0, "units_added": 0}
    with transaction.atomic():
        seen_articles = set()
        for line_number, row in enumerate(rows, start=2):
            try:
                with transaction.atomic():
                    _import_row(row, user, seen_articles, summary, apply_movement, StockMovement)
            except (ValueError, Category.DoesNotExist, AgeRange.DoesNotExist) as exc:
                message = str(exc) if isinstance(exc, ValueError) else (
                    f"catégorie « {row.get('categorie')} » introuvable"
                    if isinstance(exc, Category.DoesNotExist) else "tranche d'âge introuvable")
                errors.append({"line": line_number, "error": message})
            except BusinessError as exc:
                errors.append({"line": line_number, "error": str(exc.detail)})
        if errors or dry_run:
            transaction.set_rollback(True)
    return {"dry_run": dry_run, "success": not errors, "errors": errors, "rows": len(rows), **summary}


def _import_row(row, user, seen_articles, summary, apply_movement, StockMovement):
    for column in REQUIRED:
        if not row.get(column):
            raise ValueError(f"la colonne « {column} » est vide")

    reference = row["reference"]
    defaults = {
        "name": row["nom"],
        "category": _find_category(row["categorie"]),
        "article_type": _choice(row.get("type"), ArticleType.choices, ArticleType.CLOTHING),
        "gender": _choice(row.get("genre"), Gender.choices, Gender.UNISEX),
        "brand": row.get("marque", ""),
        "purchase_price": _decimal(row["prix_achat"], "prix_achat"),
        "sale_price": _decimal(row["prix_vente"], "prix_vente"),
    }
    article = Article.objects.filter(reference=reference).first()
    if article is None:
        article = Article.objects.create(reference=reference, created_by=user, **defaults)
        summary["articles_created"] += 1
    elif reference not in seen_articles:
        for field, value in defaults.items():
            setattr(article, field, value)
        article.save()
        summary["articles_updated"] += 1
    seen_articles.add(reference)

    if row.get("tranches_age"):
        labels = [label.strip() for label in row["tranches_age"].split("|") if label.strip()]
        for label in labels:
            article.age_ranges.add(AgeRange.objects.get(label__iexact=label))

    size = color = None
    if row.get("taille"):
        size = Size.objects.filter(label__iexact=row["taille"]).first() or Size.objects.create(
            label=row["taille"], system=Size.System.OTHER)
    if row.get("couleur"):
        color = Color.objects.filter(name__iexact=row["couleur"]).first() or Color.objects.create(
            name=row["couleur"].capitalize())

    variant = ArticleVariant.objects.filter(article=article, size=size, color=color).first()
    if variant is None:
        barcode = row.get("code_barres") or None
        if barcode and ArticleVariant.objects.filter(barcode=barcode).exists():
            raise ValueError(f"le code-barres {barcode} est déjà utilisé")
        variant = ArticleVariant.objects.create(
            article=article, size=size, color=color, barcode=barcode,
            alert_threshold=_integer(row.get("seuil_alerte"), "seuil_alerte", default=2),
        )
        summary["variants_created"] += 1

    quantity = _integer(row.get("quantite"), "quantite")
    if quantity:
        lot_number = row.get("numero_lot", "")
        if article.is_regulated and not lot_number:
            raise ValueError("numéro de lot obligatoire pour ce produit réglementé")
        apply_movement(variant, StockMovement.Type.PURCHASE, quantity, user,
                       reason="Import de fichier", lot_number=lot_number, unit_cost=article.purchase_price)
        summary["units_added"] += quantity
